"""時間割を表にして、Excel に書き出す（Googleスプレッドシートへは google_sheets.write_sheets）。"""
from __future__ import annotations

import colorsys
import datetime
import io
import itertools
import math
import unicodedata
from dataclasses import dataclass

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Problem, Result
from .previous import SHEET_PREVIOUS, previous_sheet

SHEET_SCHEDULE = "スケジュール表"
SHEET_ALL = "時間割_全体"
SHEET_STUDENT = "時間割_生徒別"
SHEET_TEACHER = "時間割_講師別"
SHEET_TEACHER_COLUMNS = "時間割_講師列"
SPARE = "予備"
SHEET_UNASSIGNED = "未配置"


# 科目の色（薄い色。文字は黒で読める）。足りなければ色相をずらして作る
PALETTE = ["#dbeafe", "#fee2e2", "#dcfce7", "#fef9c3", "#f3e8ff", "#ffedd5", "#ccfbf1", "#fce7f3",
           "#e0e7ff", "#ecfccb", "#cffafe", "#fae8ff", "#e7e5e4"]


def subject_colors(subjects: list[str]) -> dict[str, str]:
    """科目ごとに色を割り当てる。並び順で決まるので、同じシートなら毎回同じ色になる。"""
    colors = {}
    for i, subj in enumerate(subjects):
        if i < len(PALETTE):
            colors[subj] = PALETTE[i]
        else:
            hue = (i * 137.5) % 360  # 黄金角で色相をずらし、隣の色と似ないようにする
            colors[subj] = f"hsl({hue:.0f}, 80%, 90%)"
    return colors


@dataclass(frozen=True)
class View:
    """表示の切り替え。画面と書き出しの両方で使う。"""
    show_room: bool = True
    show_detail: bool = True  # 生徒の学年と、生徒・講師のIDを出す


def student_name(problem: Problem, sid: str, view: View = View()) -> str:
    st = problem.students[sid]
    return st.label if view.show_detail else st.name


def _students(problem: Problem, sids, view: View) -> str:
    return "、".join(student_name(problem, x, view) for x in sids)


def _room(room: str, view: View) -> str:
    return f" / {room}" if view.show_room else ""


def _koma_label(problem: Problem, no: int) -> str:
    k = next(k for k in problem.settings.komas if k.no == no)
    return f"{no}限 {k.start}-{k.end}"


def lessons_table(problem: Problem, result: Result, view: View = View()) -> pd.DataFrame:
    rows = []
    for ls in result.lessons:
        rows.append({
            "曜日": ls.slot.day,
            "コマ": _koma_label(problem, ls.slot.koma),
            "教室": ls.room,
            "形式": ls.form,
            "科目": ls.subject,
            "クラス": ls.class_name,
            "講師": problem.teachers[ls.teacher_id].name,
            "生徒": _students(problem, ls.student_ids, view),
            "人数": len(ls.student_ids),
        })
    columns = ["曜日", "コマ", "教室", "形式", "科目", "クラス", "講師", "生徒", "人数"]
    if not view.show_room:
        columns.remove("教室")
    return pd.DataFrame(rows, columns=columns)


def overview_grid(problem: Problem, result: Result, view: View = View()) -> pd.DataFrame:
    """行 = コマ、列 = 曜日 の全体表。セルにはその時間の授業を並べる。"""
    s = problem.settings
    grid = pd.DataFrame("", index=[_koma_label(problem, k.no) for k in s.komas], columns=s.days)
    for ls in result.lessons:
        t = problem.teachers[ls.teacher_id].name
        room = f"[{ls.room}] " if view.show_room else ""
        text = (f"{room}{ls.class_name}（{t}）{len(ls.student_ids)}名" if ls.form == "集団"
                else f"{room}個別 {ls.subject}（{t}）{_students(problem, ls.student_ids, view)}")
        cell = grid.at[_koma_label(problem, ls.slot.koma), ls.slot.day]
        grid.at[_koma_label(problem, ls.slot.koma), ls.slot.day] = (cell + "\n" if cell else "") + text
    return grid


def person_grid(problem: Problem, result: Result, kind: str, pid: str, view: View = View()) -> pd.DataFrame:
    """1人分の週間時間割。kind は "student" か "teacher"。"""
    s = problem.settings
    grid = pd.DataFrame("", index=[_koma_label(problem, k.no) for k in s.komas], columns=s.days)
    person = problem.students[pid] if kind == "student" else problem.teachers[pid]
    for sl in s.slots:
        if sl not in person.available:
            grid.at[_koma_label(problem, sl.koma), sl.day] = "－"
    for ls in result.lessons:
        if (kind == "student" and pid in ls.student_ids) or (kind == "teacher" and ls.teacher_id == pid):
            if kind == "student":
                subj = ls.student_subjects.get(pid, ls.subject)
                text = f"{subj}（{ls.form}）{problem.teachers[ls.teacher_id].name}{_room(ls.room, view)}"
            else:
                who = ls.class_name if ls.form == "集団" else _students(problem, ls.student_ids, view)
                text = f"{ls.subject}（{ls.form}）{who}{_room(ls.room, view)}"
            grid.at[_koma_label(problem, ls.slot.koma), ls.slot.day] = text
    return grid


def teacher_columns_grid(problem: Problem, result: Result, view: View = View()) -> pd.DataFrame:
    """行 = コマ、列 = 曜日 → その日に授業がある講師 の表。

    授業のない講師の列はその曜日に出さず、各曜日の最後に空の「予備」列を1列入れる。
    """
    s = problem.settings
    index = [_koma_label(problem, k.no) for k in s.komas]
    cells: dict[tuple[str, str], dict[str, str]] = {}
    for ls in result.lessons:
        who = ls.class_name if ls.form == "集団" else _students(problem, ls.student_ids, view)
        cells.setdefault((ls.slot.day, ls.teacher_id), {})[_koma_label(problem, ls.slot.koma)] = \
            f"{ls.subject}（{ls.form}）{who}{_room(ls.room, view)}"
    columns, data = [], []
    for d in s.days:
        for tid, t in problem.teachers.items():
            if (d, tid) in cells:
                columns.append((d, t.name))
                data.append([cells[d, tid].get(i, "") for i in index])
        columns.append((d, SPARE))
        data.append([""] * len(index))
    grid = pd.DataFrame(dict(zip(range(len(columns)), data)), index=index)
    grid.columns = pd.MultiIndex.from_tuples(columns, names=["曜日", "講師"])
    grid.index.name = "コマ"
    return grid


def unassigned_table(problem: Problem, result: Result, view: View = View()) -> pd.DataFrame:
    rows = [{
        "生徒ID": u.request.student_id,
        "氏名": problem.students[u.request.student_id].name,
        "科目": u.request.subject,
        "形式": u.request.form,
        "希望回数": u.request.times,
        "入らなかった回数": u.missing,
        "理由": u.reason,
    } for u in result.unassigned]
    columns = ["生徒ID", "氏名", "科目", "形式", "希望回数", "入らなかった回数", "理由"]
    if not view.show_detail:
        columns.remove("生徒ID")
    return pd.DataFrame(rows, columns=columns)


def _stacked(problem: Problem, result: Result, kind: str, view: View) -> pd.DataFrame:
    people = problem.students if kind == "student" else problem.teachers
    frames = []
    for pid, p in people.items():
        g = person_grid(problem, result, kind, pid, view).reset_index(names="コマ")
        g.insert(0, "氏名", p.name)
        if view.show_detail:
            g.insert(0, "ID", pid)
        frames.append(g)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def all_sheets(problem: Problem, result: Result, view: View = View()) -> dict[str, pd.DataFrame]:
    return {
        SHEET_ALL: overview_grid(problem, result, view).reset_index(names="コマ"),
        SHEET_STUDENT: _stacked(problem, result, "student", view),
        SHEET_TEACHER: _stacked(problem, result, "teacher", view),
        SHEET_TEACHER_COLUMNS: teacher_columns_grid(problem, result, view),
        SHEET_UNASSIGNED: unassigned_table(problem, result, view),
        "授業一覧": lessons_table(problem, result, view),
    }


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def _time_axis(problem: Problem) -> tuple[int, int, int]:
    """スケジュール表の時間の刻み（分）・最初の行の時刻・最後の時刻を決める。

    紙の表に合わせて30分刻みにし、コマはそのコマの時間にかかる行をまとめて1つのマスにする。
    30分刻みだと隣のコマと行が重なるときは、コマの時刻が割り切れる刻みにする。
    """
    times = sorted((_minutes(k.start), _minutes(k.end)) for k in problem.settings.komas)
    step = 30
    spans = [(a // step * step, -(-b // step) * step) for a, b in times]
    if any(prev[1] > nxt[0] for prev, nxt in itertools.pairwise(spans)):
        step = math.gcd(*[t for ab in times for t in ab], 60)
        spans = times
    return step, spans[0][0], max(b for _, b in spans)


def _rgb(color: str) -> str:
    """subject_colors の色（#rrggbb か hsl(h, s%, l%)）を Excel の RRGGBB にする。"""
    if color.startswith("#"):
        return color[1:].upper()
    h, sat, light = (float(x.strip(" %")) for x in color[4:-1].split(","))
    r, g, b = colorsys.hls_to_rgb(h / 360, light / 100, sat / 100)
    return f"{round(r * 255):02X}{round(g * 255):02X}{round(b * 255):02X}"


def _schedule_text(problem: Problem, ls, view: View) -> str:
    """スケジュール表の1マス。個別は「生徒　科目」を1人1行、集団は「科目　学年」。"""
    if ls.form == "集団":
        lines = [ls.class_name.replace(" ", "\u3000")]
    else:
        lines = [f"{student_name(problem, x, view)}\u3000{ls.student_subjects.get(x, ls.subject)}"
                 for x in ls.student_ids]
    if view.show_room and ls.room:
        lines.append(f"（{ls.room}）")
    return "\n".join(lines)


def _text_lines(text: str, width: float) -> int:
    """列幅 width の Excel のマスで text が何行になるかの目安（全角は2文字ぶん）。"""
    return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)
                                / (width - 2))) for line in text.split("\n"))


def write_schedule_sheet(ws, problem: Problem, result: Result, view: View = View()) -> None:
    """紙のスケジュール表と同じ形で書く。

    曜日ごとに「時刻の列＋その日に授業がある講師の列＋予備の列」を横に並べ、
    行は30分刻みの時刻、授業はコマの時間ぶんの行をまとめた1マスに書く。
    上に曜日と講師名、下にも講師名と曜日を入れる。
    """
    step, first, last = _time_axis(problem)
    n_rows = (last - first) // step
    top = 3                    # 1行目：曜日、2行目：講師名、3行目から時刻
    foot = top + n_rows        # 下の講師名の行（その下に曜日）

    font = "游ゴシック"
    medium, thin, hair = (Side(style=x) for x in ("medium", "thin", "hair"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    colors = {k: _rgb(v) for k, v in subject_colors(problem.subjects).items()}

    width = 22
    row_height = 20.0  # 時刻の1行の高さ。授業の文字がマスに収まるよう、下で広げる
    by_cell = {(ls.slot.day, ls.teacher_id, ls.slot.koma): ls for ls in result.lessons}
    komas = {k.no: k for k in problem.settings.komas}

    col = 1
    for d in problem.settings.days:
        teachers = [(tid, t.name) for tid, t in problem.teachers.items()
                    if any(ls.slot.day == d and ls.teacher_id == tid for ls in result.lessons)]
        teachers.append((None, SPARE))
        left, right = col, col + len(teachers)

        for row in (1, foot + 1):
            ws.merge_cells(start_row=row, start_column=left, end_row=row, end_column=right)
            c = ws.cell(row, left, d)
            c.font = Font(name=font, size=16, bold=True)
            c.alignment = center
        ws.column_dimensions[get_column_letter(left)].width = 8
        for i in range(n_rows):
            minute = first + i * step
            c = ws.cell(top + i, left, datetime.time(minute // 60, minute % 60))
            c.number_format = "h:mm"
            c.font = Font(name=font, size=11)
            c.alignment = Alignment(horizontal="right", vertical="top")

        for j, (tid, name) in enumerate(teachers, start=left + 1):
            ws.column_dimensions[get_column_letter(j)].width = width
            for row in (top - 1, foot):
                c = ws.cell(row, j, name)
                c.font = Font(name=font, size=12, bold=True)
                c.alignment = center
            for no, k in komas.items():
                ls = by_cell.get((d, tid, no))
                if ls is None:
                    continue
                r0 = top + (_minutes(k.start) - first) // step
                r1 = top + -(-(_minutes(k.end) - first) // step) - 1
                if r1 > r0:
                    ws.merge_cells(start_row=r0, start_column=j, end_row=r1, end_column=j)
                c = ws.cell(r0, j, _schedule_text(problem, ls, view))
                c.font = Font(name=font, size=12, bold=True)
                row_height = max(row_height, (_text_lines(c.value, width) * 16 + 4) / (r1 - r0 + 1))
                c.alignment = center
                subj = ls.student_subjects.get(ls.student_ids[0], ls.subject) if ls.student_ids else ls.subject
                if subj in colors:
                    c.fill = PatternFill("solid", fgColor=colors[subj])

        # 罫線：曜日のまとまりは太線、講師の間は細線、時刻の行は極細線、授業のマスは細線で囲む
        lesson_rows = {(m.min_col, r): (m.min_row, m.max_row) for m in ws.merged_cells.ranges
                       if top <= m.min_row < foot for r in range(m.min_row, m.max_row + 1)}
        for j in range(left, right + 1):
            for row in range(1, foot + 2):
                first_row, last_row = lesson_rows.get((j, row), (None, None))
                if row < top or row >= foot - 1:
                    bottom = medium
                elif last_row is not None:
                    bottom = thin if row == last_row else None
                else:
                    bottom = hair
                ws.cell(row, j).border = Border(
                    left=medium if j == left else thin,
                    right=medium if j == right else thin,
                    top=medium if row in (1, foot) else thin if row == first_row else None,
                    bottom=bottom,
                )
        col = right + 1

    ws.row_dimensions[1].height = ws.row_dimensions[foot + 1].height = 28
    ws.row_dimensions[top - 1].height = ws.row_dimensions[foot].height = 22
    for row in range(top, foot):
        ws.row_dimensions[row].height = math.ceil(row_height)
    ws.freeze_panes = ws.cell(top, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def to_xlsx(problem: Problem, result: Result, view: View = View()) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        # 最初のシートは紙のスケジュール表と同じ形（講師列の表を、時刻の行とまとめたマスで書いたもの）
        write_schedule_sheet(xw.book.create_sheet(SHEET_SCHEDULE), problem, result, view)
        for name, df in all_sheets(problem, result, view).items():
            if name == SHEET_TEACHER_COLUMNS:
                continue
            # 曜日と講師の2段の見出しは、コマを行の見出しにして書く（Excel で曜日がまとまる）
            df.to_excel(xw, sheet_name=name, index=isinstance(df.columns, pd.MultiIndex))
            ws = xw.sheets[name]
            for col in ws.iter_cols():
                width = max((len(str(c.value or "")) for c in col), default=4)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(8, width * 1.6))
        # 次の作り直しで「今の時間割をなるべく変えない」ために読む記録（表示の切り替えに関係なくIDで持つ）
        previous_sheet(result).to_excel(xw, sheet_name=SHEET_PREVIOUS, index=False)
        xw.sheets[SHEET_PREVIOUS].sheet_state = "hidden"
    return buf.getvalue()
