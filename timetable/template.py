"""入力用Excelの雛形を作る。見出しだけが入っていて、そのまま書き込んでアップロードできる。"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .loader import SHEET_KOMA, SHEET_STUDENT_AVAIL, SHEET_STUDENT_REQ, SHEET_TEACHER, SHEET_TEACHER_AVAIL
from .model import SUBJECTS, Settings

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def template_sheets(settings: Settings) -> dict[str, pd.DataFrame]:
    """雛形の各シート。空き時間シートの列は今のコマ割（月1, 月2 …）に合わせる。"""
    slots = [s.label for s in settings.slots]
    return {
        SHEET_KOMA: pd.DataFrame([{"コマ": k.no, "開始": k.start, "終了": k.end} for k in settings.komas]),
        SHEET_STUDENT_REQ: pd.DataFrame(columns=["生徒ID", "氏名", "学年", "科目", "形式", "週の回数", "希望講師（任意）"]),
        SHEET_STUDENT_AVAIL: pd.DataFrame(columns=["生徒ID", "氏名", *slots]),
        SHEET_TEACHER: pd.DataFrame(columns=["講師ID", "氏名", *SUBJECTS, "集団可", "個別可", "週の上限コマ"]),
        SHEET_TEACHER_AVAIL: pd.DataFrame(columns=["講師ID", "氏名", *slots]),
    }


def template_xlsx(settings: Settings) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in template_sheets(settings).items():
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "C2" if "氏名" in df.columns else "A2"
            for i, col in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
                ws.column_dimensions[get_column_letter(i)].width = max(6, len(str(col)) * 2 + 2)
    return buf.getvalue()
