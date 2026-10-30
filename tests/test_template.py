"""入力用の雛形が、今のコマ割に合っていて、そのまま読み込めることを確かめる。"""
import io
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from timetable.loader import REQUIRED_SHEETS, load_settings, parse, read_xlsx  # noqa: E402
from timetable.template import EXAMPLE_SUBJECTS as SUBJECTS, make_template  # noqa: E402


def test_template_matches_settings():
    settings = load_settings(ROOT / "config.toml")
    settings.days = ["月", "水"]
    wb = openpyxl.load_workbook(io.BytesIO(make_template(settings)))
    assert set(REQUIRED_SHEETS) <= set(wb.sheetnames)
    header = [c.value for c in wb["生徒_空き時間"][1]]
    assert header == ["生徒ID", "氏名", *[s.label for s in settings.slots]]
    assert wb["生徒_空き時間"].data_validations.dataValidation


def test_filled_template_can_be_read():
    settings = load_settings(ROOT / "config.toml")
    wb = openpyxl.load_workbook(io.BytesIO(make_template(settings)))
    wb["生徒_受講希望"].append(["S1", "山田", "高1", "数学", "個別", 1, ""])
    wb["生徒_空き時間"].append(["S1", "山田", "○"])
    wb["講師_担当可能科目"].append(["T1", "佐藤", *["○" if s == "数学" else "" for s in SUBJECTS], "", "○", ""])
    wb["講師_空き時間"].append(["T1", "佐藤", "○"])
    buf = io.BytesIO()
    wb.save(buf)
    problem, notes = parse(read_xlsx(buf.getvalue()), settings)
    assert problem is not None, notes
    assert problem.students["S1"].available == {settings.slots[0]}
    assert problem.teachers["T1"].subjects == {"数学"}


def test_template_keeps_people(tmp_path):
    settings = load_settings(ROOT / "config.toml")
    problem, _ = parse(read_xlsx(ROOT / "sample" / "サンプル入力.xlsx"), settings)
    wb = openpyxl.load_workbook(io.BytesIO(make_template(settings, problem)))
    ids = [r[0].value for r in wb["講師_空き時間"].iter_rows(min_row=2)]
    assert ids == list(problem.teachers)
    assert all(c.value is None for r in wb["講師_空き時間"].iter_rows(min_row=2, min_col=3) for c in r)


def test_template_uses_subjects_from_input():
    settings = load_settings(ROOT / "config.toml")
    problem, _ = parse(read_xlsx(ROOT / "sample" / "サンプル入力.xlsx"), settings)
    problem.subjects = ["英語", "英検対策"]
    wb = openpyxl.load_workbook(io.BytesIO(make_template(settings, problem)))
    assert [c.value for c in wb["講師_担当可能科目"][1]][2:4] == ["英語", "英検対策"]
