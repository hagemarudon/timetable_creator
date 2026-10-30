"""サンプル入力で時間割を組み、必須条件が守られているかを確かめる。"""
import subprocess
import sys
from collections import Counter
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from timetable.loader import load_settings, parse, read_xlsx  # noqa: E402
from timetable.solver import solve  # noqa: E402


@pytest.fixture(scope="module")
def solved(tmp_path_factory):
    xlsx = tmp_path_factory.mktemp("in") / "sample.xlsx"
    subprocess.run([sys.executable, "make_sample.py", str(xlsx)], cwd=ROOT, check=True)
    settings = load_settings(ROOT / "config.toml")
    settings.time_limit_seconds = 20
    problem, _ = parse(read_xlsx(xlsx), settings)
    return problem, solve(problem)


def test_solved(solved):
    _, result = solved
    assert result.status in ("最適解", "時間内の最良解")
    assert result.lessons


def test_people_only_in_available_slots(solved):
    problem, result = solved
    for ls in result.lessons:
        assert ls.slot in problem.teachers[ls.teacher_id].available
        for sid in ls.student_ids:
            assert ls.slot in problem.students[sid].available


def test_no_double_booking(solved):
    problem, result = solved
    t = Counter((ls.teacher_id, ls.slot) for ls in result.lessons)
    assert max(t.values()) == 1
    s = Counter((sid, ls.slot) for ls in result.lessons for sid in ls.student_ids)
    assert max(s.values()) == 1


def test_teacher_can_teach(solved):
    problem, result = solved
    for ls in result.lessons:
        t = problem.teachers[ls.teacher_id]
        assert set(ls.subject.split("・")) <= t.subjects
        assert t.can_group if ls.form == "集団" else t.can_individual


def test_capacity_and_rooms(solved):
    problem, result = solved
    s = problem.settings
    for ls in result.lessons:
        limit = s.group_capacity if ls.form == "集団" else s.max_students_per_teacher
        assert len(ls.student_ids) <= limit
    per_slot = Counter((ls.slot, ls.form) for ls in result.lessons)
    for (sl, form), n in per_slot.items():
        assert n <= (s.group_rooms if form == "集団" else s.individual_booths)


def test_times_and_teacher_limit(solved):
    problem, result = solved
    got = Counter()
    for ls in result.lessons:
        for sid in ls.student_ids:
            r = next(r for r in problem.requests
                     if r.student_id == sid and r.form == ls.form and r.subject == ls.student_subjects[sid])
            got[r.index] += 1
    for r in problem.requests:
        missing = next((u.missing for u in result.unassigned if u.request.index == r.index), 0)
        assert got[r.index] == r.times - missing
    load = Counter(ls.teacher_id for ls in result.lessons)
    for tid, n in load.items():
        mx = problem.teachers[tid].max_komas
        assert mx is None or n <= mx


def test_same_name_students_are_distinguishable(solved):
    from timetable.output import lessons_table

    problem, result = solved
    names = [s.name for s in problem.students.values()]
    labels = [s.label for s in problem.students.values()]
    assert len(set(labels)) == len(labels)
    for s in problem.students.values():
        assert s.id in s.label and s.grade in s.label
    table = lessons_table(problem, result)
    assert table["生徒"].str.contains("（").all() or table.empty
    assert names  # 氏名そのものは変えない


def _koma_runs_ok(komas: list[int]) -> bool:
    return sorted(komas) == list(range(min(komas), max(komas) + 1))


def test_lessons_are_consecutive_for_students_and_teachers(solved):
    problem, result = solved
    days: dict[tuple[str, str, str], list[int]] = {}
    for ls in result.lessons:
        days.setdefault(("t", ls.teacher_id, ls.slot.day), []).append(ls.slot.koma)
        for sid in ls.student_ids:
            days.setdefault(("s", sid, ls.slot.day), []).append(ls.slot.koma)
    order = {k.no: i for i, k in enumerate(problem.settings.komas)}
    for key, komas in days.items():
        assert _koma_runs_ok([order[k] for k in komas]), key


def test_teacher_columns_grid(solved):
    from timetable.output import SPARE, teacher_columns_grid

    problem, result = solved
    grid = teacher_columns_grid(problem, result)
    assert list(grid.index) == [f"{k.no}限 {k.start}-{k.end}" for k in problem.settings.komas]
    assert list(dict.fromkeys(d for d, _ in grid.columns)) == problem.settings.days
    for d in problem.settings.days:
        cols = [t for day, t in grid.columns if day == d]
        assert cols[-1] == SPARE and cols.count(SPARE) == 1
        assert (grid[(d, SPARE)] == "").all()
        working = {problem.teachers[ls.teacher_id].name for ls in result.lessons if ls.slot.day == d}
        assert set(cols[:-1]) == working
        for t in cols[:-1]:
            assert (grid[(d, t)] != "").any()


def test_teacher_columns_grid_day_without_lessons(solved):
    from timetable.model import Result
    from timetable.output import SPARE, teacher_columns_grid, to_xlsx

    problem, _ = solved
    grid = teacher_columns_grid(problem, Result("最適解", [], []))
    assert list(grid.columns) == [(d, SPARE) for d in problem.settings.days]
    assert to_xlsx(problem, Result("最適解", [], []))


def test_xlsx_schedule_sheet_like_paper(solved):
    """最初のシートは紙のスケジュール表の形：曜日→時刻＋講師列、30分刻みの行、授業はコマの行をまとめたマス。"""
    import datetime
    import io

    import openpyxl

    from timetable.output import SHEET_SCHEDULE, SHEET_TEACHER_COLUMNS, to_xlsx

    problem, result = solved
    wb = openpyxl.load_workbook(io.BytesIO(to_xlsx(problem, result)))
    assert wb.sheetnames[0] == SHEET_SCHEDULE and SHEET_TEACHER_COLUMNS not in wb.sheetnames
    ws = wb[SHEET_SCHEDULE]
    days = problem.settings.days
    assert ws.cell(1, 1).value == days[0]
    assert ws.cell(3, 1).value == datetime.time(16, 0)
    assert ws.cell(4, 1).value == datetime.time(16, 30)
    assert [c.value for c in ws[1] if c.value] == days
    assert [c.value for c in ws[ws.max_row] if c.value] == days
    assert [c.value for c in ws[2]].count("予備") == len(days)
    assert [c.value for c in ws[2]] == [c.value for c in ws[ws.max_row - 1]]

    # 授業の数だけ、3行（80分のコマ）をまとめたマスがある
    lesson_merges = [m for m in ws.merged_cells.ranges if m.min_row >= 3 and m.min_col == m.max_col]
    assert len(lesson_merges) == len(result.lessons)
    assert all(m.max_row - m.min_row == 2 for m in lesson_merges)
    ls = next(x for x in result.lessons if x.form == "個別")
    text = "\n".join(str(ws.cell(m.min_row, m.min_col).value) for m in lesson_merges)
    assert f"{problem.students[ls.student_ids[0]].name}" in text


def test_schedule_time_axis_falls_back_when_komas_overlap(solved):
    from timetable.model import Koma
    from timetable.output import _rgb, _time_axis

    problem, _ = solved
    assert _time_axis(problem) == (30, 16 * 60, 22 * 60)
    komas = problem.settings.komas
    problem.settings.komas = [Koma(1, "16:00", "17:10"), Koma(2, "17:10", "18:20")]
    try:
        assert _time_axis(problem) == (10, 16 * 60, 18 * 60 + 20)
    finally:
        problem.settings.komas = komas
    assert _rgb("#dbeafe") == "DBEAFE"
    assert len(_rgb("hsl(200, 80%, 90%)")) == 6


def test_view_hides_room_and_student_detail(solved):
    from timetable.output import View, all_sheets, lessons_table, teacher_columns_grid

    problem, result = solved
    shown = lessons_table(problem, result)
    assert "教室" in shown.columns
    sids = [sid for ls in result.lessons for sid in ls.student_ids]
    st = problem.students[sids[0]]
    assert st.id in "".join(shown["生徒"])

    view = View(show_room=False, show_detail=False)
    hidden = lessons_table(problem, result, view)
    assert "教室" not in hidden.columns
    text = "".join(hidden["生徒"]) + teacher_columns_grid(problem, result, view).to_string()
    rooms = {ls.room for ls in result.lessons}
    assert not any(r in text for r in rooms)
    assert not any(problem.students[x].id in text for x in sids)
    sheets = all_sheets(problem, result, view)
    assert "ID" not in sheets["時間割_生徒別"].columns
    assert "ID" not in sheets["時間割_講師別"].columns
    assert "ID" in all_sheets(problem, result)["時間割_講師別"].columns
